from datetime import datetime, date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from core.constants import REPORTS_DIR
from core.database import Database


class ExcelManager:
    """Excel se genera bajo demanda (o al cierre de mes), no en cada heartbeat."""

    HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
    HEADER_FONT = Font(bold=True, color="FFFFFF")

    @staticmethod
    def _month_file(month: date = None) -> Path:
        m = (month or date.today()).strftime("%Y-%m")
        return REPORTS_DIR / f"{m}.xlsx"

    @classmethod
    def _style_header(cls, ws):
        for cell in ws[1]:
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

    @classmethod
    def export_month(cls, month: date = None) -> Path:
        """Genera el xlsx del mes indicado (por defecto, mes actual)."""
        month = month or date.today()
        first = month.replace(day=1)
        if month.month == 12:
            next_first = month.replace(year=month.year + 1, month=1, day=1)
        else:
            next_first = month.replace(month=month.month + 1, day=1)

        events = Database.events_between(
            datetime.combine(first, datetime.min.time()),
            datetime.combine(next_first, datetime.min.time()),
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "Eventos"
        ws.append(["FechaHora", "Evento", "Detalle"])
        cls._style_header(ws)

        for ev in events:
            ws.append([ev["timestamp"], ev["event"], ev["detail"]])

        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 50

        # Hoja Resumen
        ws2 = wb.create_sheet("Resumen")
        ws2.append(["Fecha", "Horas_ON", "Horas_OFF", "Reinicios", "Cortes"])
        cls._style_header(ws2)

        with Database._conn() as c:
            rows = c.execute(
                "SELECT * FROM summaries WHERE day BETWEEN ? AND ? "
                "ORDER BY day",
                (first.isoformat(), next_first.isoformat()),
            ).fetchall()

        for r in rows:
            ws2.append([
                r["day"], r["hours_on"], r["hours_off"],
                r["reboots"], r["cuts"],
            ])

        for col in "ABCDE":
            ws2.column_dimensions[col].width = 14

        out = cls._month_file(month)
        wb.save(out)
        return out

    @classmethod
    def export_all(cls):
        """Exporta todos los meses que tengan eventos."""
        with Database._conn() as c:
            rows = c.execute(
                "SELECT DISTINCT substr(timestamp,1,7) AS m FROM events ORDER BY m"
            ).fetchall()

        generated = []
        for r in rows:
            y, m = map(int, r["m"].split("-"))
            generated.append(cls.export_month(date(y, m, 1)))
        return generated